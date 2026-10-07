from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.app.utils.text import sanitize_text


ASSET_DIR = Path(__file__).resolve().parents[3] / "assets"
LOGO_PATH = ASSET_DIR / "logo.png"


def _set_cell_shading(cell, fill: str = "E8EEF7"):
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def _set_default_font(document: Document):
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    # Ensure Word uses Times New Roman for East Asian font mappings too.
    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "Times New Roman",
    )


def _add_footer(document: Document):
    for section in document.sections:
        footer = section.footer

        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            "Generated with LegalEase — AI-assisted legal document draft"
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(8)


def _extract_terms(text: str):
    """
    Find common terms/conditions represented as semicolon-separated
    clauses in the original user input when they are still present
    in the generated content.
    """

    matches = re.findall(
        r"(?:^|\n)\s*(?:[-•]\s*)?([^;\n]{10,300});",
        text,
    )

    return [match.strip() for match in matches[:20]]


def format_docx(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)

    document = Document()

    _set_default_font(document)

    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # Logo
    if LOGO_PATH.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()
        run.add_picture(
            str(LOGO_PATH),
            width=Inches(1.25),
        )

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(doc_type.upper())
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(16)

    document.add_paragraph()

    # Main document
    for block in text.split("\n\n"):
        block = block.strip()

        if not block:
            continue

        lines = block.splitlines()

        for line in lines:
            line = line.strip()

            if not line:
                continue

            if re.match(r"^\d+\.\s+", line):
                paragraph = document.add_paragraph()

                run = paragraph.add_run(line)
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            elif line.startswith(("-", "•", "*")):
                paragraph = document.add_paragraph(
                    style="List Bullet"
                )

                run = paragraph.add_run(
                    re.sub(r"^[-•*]\s*", "", line)
                )

                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

            else:
                paragraph = document.add_paragraph()

                run = paragraph.add_run(line)
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

                paragraph.paragraph_format.space_after = Pt(7)
                paragraph.paragraph_format.line_spacing = 1.15

    # Terms table
    terms = _extract_terms(text)

    if terms:
        document.add_page_break()

        heading = document.add_paragraph()
        run = heading.add_run("Key Terms")
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)

        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.style = "Table Grid"

        header = table.rows[0].cells

        header[0].text = "No."
        header[1].text = "Term / Condition"

        _set_cell_shading(header[0])
        _set_cell_shading(header[1])

        for index, term in enumerate(terms, start=1):
            cells = table.add_row().cells

            cells[0].text = str(index)
            cells[1].text = term

    _add_footer(document)

    output = BytesIO()
    document.save(output)

    return output.getvalue()